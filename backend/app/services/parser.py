from dataclasses import dataclass, field
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

WORD_THRESHOLD = 100
CHUNK_THRESHOLD = 6000

@dataclass
class ParsedDocument:
    text: str
    comments: list[str]
    image_paths: list[str]
    word_count: int
    too_short: bool
    needs_chunking: bool

def parse_docx(path: Path) -> ParsedDocument:
    try:
        doc = Document(str(path))
    except Exception:
        return _fallback_mammoth(path)

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)
    comments = _extract_comments(doc)
    if comments:
        text += "\n\n[AUTHOR COMMENTS]\n" + "\n".join(comments)

    image_paths = [
        rel.target_ref
        for rel in doc.part.rels.values()
        if "image" in rel.reltype
    ]

    wc = len(text.split())
    return ParsedDocument(
        text=text,
        comments=comments,
        image_paths=image_paths,
        word_count=wc,
        too_short=wc < WORD_THRESHOLD,
        needs_chunking=wc > CHUNK_THRESHOLD,
    )

def _extract_comments(doc: Document) -> list[str]:
    try:
        comments_part = doc.part.package.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
        )
        root = comments_part._element
        return [
            "".join(r.text for r in c.findall(f".//{qn('w:r')}/{qn('w:t')}"))
            for c in root.findall(qn("w:comment"))
        ]
    except Exception:
        return []

def _fallback_mammoth(path: Path) -> ParsedDocument:
    import mammoth
    with open(path, "rb") as f:
        result = mammoth.extract_raw_text(f)
    text = result.value
    wc = len(text.split())
    return ParsedDocument(
        text=text, comments=[], image_paths=[],
        word_count=wc,
        too_short=wc < WORD_THRESHOLD,
        needs_chunking=wc > CHUNK_THRESHOLD,
    )
